"""Editable, template-neutral Word export from résumé JSON."""
import io
import uuid

import docx
from docx.shared import Emu

from backend.models.db import Job, Resume, Setting

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DATA = {
    "header": {"name": "Dana Okonkwo", "title": "Senior Product Manager",
               "contact_items": [{"text": "dana@example.com", "url": "mailto:dana@example.com"},
                                 {"text": "Berlin", "url": ""}]},
    "summary": "Product leader.",
    "experience": [{"title": "PM", "company": "Acme", "location": "Berlin", "date": "2020 - 2024",
                    "description": "Ran the platform.", "bullets": ["Shipped **Go** services", "Cut cost 20%"]}],
    "skills": {"Languages": "Go, Python"},
    "education": [{"school": "TU Berlin", "location": "Berlin", "degree": "BSc CS"}],
    "projects": [{"name": "Side", "description": "A thing", "bullets": ["Did it"]}],
    "publications": [],
}

PNG = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="


def _resume(test_db, image=False, data=None, **over):
    test_db.add(Setting(key="dashboard_api_key", value=""))
    if image:
        test_db.add(Setting(key="profile_image_path", value=PNG))
    over.setdefault("template", "inter")
    over.setdefault("is_base", True)
    resume = Resume(name="Base CV", page_format="letter", json_data=data or DATA, **over)
    test_db.add(resume)
    test_db.commit()
    return resume


def _fetch(api_client, resume, query=""):
    response = api_client.get(f"/api/resumes/{resume.id}/docx{query}")
    assert response.status_code == 200, response.text
    return response, docx.Document(io.BytesIO(response.content))


def test_docx_download_headers(api_client, test_db):
    response, _ = _fetch(api_client, _resume(test_db))
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.content[:2] == b"PK"
    assert 'filename="DanaOkonkwo_BaseCV_Resume.docx"' in response.headers["content-disposition"]


def test_docx_repeats_tailored_resume_footer_on_each_page(api_client, test_db):
    job = Job(external_id="docx-footer", content_hash="docx-footer", title="Staff Engineer", company="Acme")
    test_db.add(job)
    test_db.flush()
    resume = _resume(test_db, is_base=False, job_id=job.id)

    _, document = _fetch(api_client, resume)

    assert document.sections[0].footer.paragraphs[0].text == "Dana Okonkwo - Staff Engineer - Acme"


def test_docx_omits_disabled_footer(api_client, test_db):
    job = Job(external_id="docx-no-footer", content_hash="docx-no-footer", title="Staff Engineer", company="Acme")
    test_db.add(job)
    test_db.flush()
    data = {**DATA, "footer_enabled": False}
    resume = _resume(test_db, data=data, is_base=False, job_id=job.id)

    _, document = _fetch(api_client, resume)

    assert document.sections[0].footer.paragraphs[0].text == ""


def test_docx_contains_every_section_in_order(api_client, test_db):
    _, document = _fetch(api_client, _resume(test_db))
    text = [paragraph.text for paragraph in document.paragraphs]
    assert text[0] == "Dana Okonkwo"
    for heading in ("Experience", "Skills", "Education", "Projects"):
        assert heading in text
    assert "Publications" not in text
    assert text.index("Skills") < text.index("Experience") < text.index("Education") < text.index("Projects")
    joined = "\n".join(text)
    for expected in ("Senior Product Manager", "Product leader.", "PM at Acme", "2020 - 2024", "Ran the platform.",
                     "Languages", "Go, Python", "TU Berlin", "BSc CS", "Side", "A thing"):
        assert expected in joined, expected


def test_docx_bullets_and_bold_markup(api_client, test_db):
    _, document = _fetch(api_client, _resume(test_db))
    bullets = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "List Bullet"]
    assert [paragraph.text for paragraph in bullets[:2]] == ["Shipped Go services", "Cut cost 20%"]
    assert [(run.text, bool(run.bold)) for run in bullets[0].runs] == [("Shipped ", False), ("Go", True), (" services", False)]


def test_docx_contact_link_is_clickable(api_client, test_db):
    _, document = _fetch(api_client, _resume(test_db))
    targets = [relation.target_ref for relation in document.part.rels.values() if relation.reltype.endswith("/hyperlink")]
    assert targets == ["mailto:dana@example.com"]


def test_docx_page_size_follows_format(api_client, test_db):
    resume = _resume(test_db)
    _, letter = _fetch(api_client, resume)
    _, a4 = _fetch(api_client, resume, "?format=a4")
    assert letter.sections[0].page_width == Emu(7772400)
    assert a4.sections[0].page_width == Emu(7560310)


def test_docx_unknown_resume_is_404(api_client, test_db):
    test_db.add(Setting(key="dashboard_api_key", value=""))
    test_db.commit()
    assert api_client.get(f"/api/resumes/{uuid.uuid4()}/docx").status_code == 404


def test_docx_reads_legacy_dates_key(api_client, test_db):
    legacy = {**DATA, "experience": [{"title": "PM", "company": "Acme", "dates": "1999 - 2001", "bullets": []}]}
    _, document = _fetch(api_client, _resume(test_db, data=legacy))
    assert any("1999 - 2001" in paragraph.text for paragraph in document.paragraphs)


def test_docx_does_not_embed_global_profile_image(api_client, test_db):
    _, document = _fetch(api_client, _resume(test_db, image=True, template="professional-top-left-image"))
    assert not document.inline_shapes and not document.tables
