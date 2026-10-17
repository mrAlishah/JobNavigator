"""`json_data["style"]` (page margins, profile image size, per-section font / size / line height)
overrides the two Professional templates in the PDF HTML and the Word export; every other template
ignores it, and invalid values are dropped or clamped."""
import io
import json
from pathlib import Path

import docx
import pytest
from docx.shared import Inches
from PIL import Image

from backend.api.routes_resumes import _build_docx, _render_html, _style_overrides

ROOT = Path(__file__).resolve().parents[1] / "resume_templates"
PRO = ["professional-top-left-image", "professional-top-right-image"]
PNG = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
DATA = {"header": {"name": "Dana", "title": "PM", "contact_items": [{"text": "Berlin"}]}, "summary": "sum",
        "experience": [{"title": "PM", "company": "Acme", "bullets": ["b1"]}],
        "skills": {"Languages": "Go"}, "education": [{"school": "TU", "degree": "BSc"}],
        "projects": [{"name": "Side"}], "publications": [{"name": "Paper"}]}
STYLE = {
    "page": {"top": 0.9, "right": 0.6, "bottom": 0.8, "left": 0.7},
    "image": {"width": 150, "height": 80},
    "sections": {"experience": {"font": "Georgia", "size": 11, "line_height": 1.4},
                 "header": {"font": "Calibri", "size": 24}},
}


def _html(tpl, style=None, **kw):
    data = {**DATA, **({"style": style} if style is not None else {})}
    return _render_html(data, tpl, "letter", profile_image_path=PNG, profile_image_enabled=True, **kw)


def _docx_style(tpl="professional-top-left-image"):
    return json.loads((ROOT / tpl / "meta.json").read_text())["docx"]


def _docx(style, tpl="professional-top-left-image", image=None, st=None):
    st = _docx_style(tpl) if st is None else st
    return docx.Document(io.BytesIO(_build_docx({**DATA, "style": style}, "letter", st, image)))


def _png(w, h):
    buf = io.BytesIO()
    Image.new("RGB", (w, h), "red").save(buf, "PNG")
    return buf.getvalue()


def test_overrides_are_validated_and_clamped():
    ov = _style_overrides({"style": {
        "page": {"top": 99, "left": "x", "right": True, "bottom": -1},
        "image": {"width": 5, "height": 9999},
        "sections": {"experience": {"font": "Comic; } body{display:none", "size": 99, "line_height": "1.4"},
                     "summary": {"font": "Georgia", "size": 2, "line_height": 0.1},
                     "bogus": {"font": "Georgia"}},
    }})
    assert ov["page"] == {"top": 1.5, "bottom": 0}
    assert ov["image"] == {"width": 40, "height": 300}
    assert ov["sections"] == {"experience": {"size": 36},
                              "summary": {"font": "Georgia", "size": 6, "line_height": 0.9}}


@pytest.mark.parametrize("junk", [None, "x", [], {"style": 5}, {"style": {"page": 5, "sections": []}}])
def test_junk_style_is_empty(junk):
    assert _style_overrides({"style": junk} if not isinstance(junk, dict) else junk) == {}


@pytest.mark.parametrize("tpl", PRO)
def test_pdf_html_applies_overrides(tpl):
    html = _html(tpl, STYLE, footer_reserved_in=0.4)
    assert ".sec-experience" in html and "Georgia" in html and "11pt" in html and "line-height: 1.4" in html
    assert "margin: 0.9in 0 0.8in 0" in html            # @page: top, bottom
    assert "padding-left: 0.7in" in html and "padding-right: 0.6in" in html
    assert "width: 150px" in html and "height: 80px" in html
    assert 'class="section sec-experience"' in html


@pytest.mark.parametrize("tpl", PRO)
def test_pdf_footer_band_wins_over_a_smaller_bottom_margin(tpl):
    assert "margin: 0.9in 0 1in 0" in _html(tpl, STYLE, footer_reserved_in=1)


@pytest.mark.parametrize("tpl", PRO)
def test_no_style_leaves_the_html_unchanged(tpl):
    assert _html(tpl, {}) == _html(tpl)


def test_other_templates_ignore_style():
    other = [p.name for p in ROOT.iterdir() if p.is_dir() and p.name not in PRO]
    for tpl in other:
        assert _html(tpl, STYLE) == _html(tpl), tpl


def test_docx_margins_fonts_sizes_and_line_spacing():
    d = _docx(STYLE)
    s = d.sections[0]
    assert (s.top_margin, s.bottom_margin) == (Inches(0.9), Inches(0.8))
    assert (s.left_margin, s.right_margin) == (Inches(0.7), Inches(0.6))
    bullet = next(p for p in d.paragraphs if p.text == "• b1")
    assert bullet.runs[0].font.name == "Georgia" and bullet.runs[0].font.size.pt == 11
    assert bullet.paragraph_format.line_spacing == 1.4
    skills = next(p for p in d.paragraphs if p.text.startswith("Languages"))
    assert skills.runs[0].font.name != "Georgia"
    name = next(p for p in d.paragraphs if p.text == "Dana")
    assert name.runs[0].font.name == "Calibri" and name.runs[0].font.size.pt == 24


def test_docx_section_title_keeps_its_own_size():
    d = _docx(STYLE)
    title = next(p for p in d.paragraphs if p.text == "PROFESSIONAL EXPERIENCE")
    assert title.runs[0].font.size.pt == _docx_style()["section_pt"]
    assert title.runs[0].font.name == "Georgia"


def test_docx_image_is_resized_and_cropped_to_the_aspect():
    d = _docx(STYLE, image=_png(200, 200))
    shape = d.inline_shapes[0]
    assert abs(shape.width - Inches(150 / 96)) <= 1 and abs(shape.height - Inches(80 / 96)) <= 1
    blob = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    img = Image.open(io.BytesIO(d.part.related_parts[blob].blob))
    assert img.size[0] / img.size[1] == pytest.approx(150 / 80, rel=0.02)


def test_docx_without_style_or_without_opt_in_is_unchanged():
    plain = _docx_style()
    off = {k: v for k, v in plain.items() if k != "style_overrides"}
    base = docx.Document(io.BytesIO(_build_docx(DATA, "letter", plain)))
    ignored = _docx(STYLE, st=off)
    assert [p.text for p in base.paragraphs] == [p.text for p in ignored.paragraphs]
    assert abs(ignored.sections[0].top_margin - Inches(plain["margin_in"][0])) < 635    # Word stores twips


@pytest.mark.parametrize("tpl", PRO)
def test_templates_mark_every_selectable_area(tpl):
    html = _html(tpl)
    for area in ("page", "header", "image", "summary", "experience", "skills", "education", "projects", "publications"):
        assert f'data-section="{area}"' in html, area


@pytest.mark.parametrize("tpl", PRO)
def test_select_mode_injects_the_click_script_only_when_asked(tpl):
    assert "resume-select" not in _html(tpl)
    html = _html(tpl, select_mode=True)
    assert html.count("resume-select") >= 1 and html.index("resume-select") < html.rindex("</body>")


def test_preview_endpoint_accepts_template_format_and_select():
    import inspect
    from backend.api.routes_resumes import preview_resume
    assert {"template", "format", "select"} <= set(inspect.signature(preview_resume).parameters)


# ── parts: every class of element in the page, with the full property set ──────────────────────────
from bs4 import BeautifulSoup
from backend.api.routes_resumes import _STYLE_PARTS, _style_css

FULL = {"header": {"name": "Dana", "title": "PM", "contact_items": [{"text": "d@x.io", "url": "mailto:d@x.io"}]}, "summary": "sum",
        "experience": [{"title": "PM", "company": "Acme", "location": "Berlin", "date": "2020", "description": "expdesc", "bullets": ["b1"]}],
        "skills": {"Languages": "Go"}, "education": [{"school": "TU", "location": "Berlin", "degree": "BSc"}],
        "projects": [{"name": "Side", "url": "x.io", "description": "d", "bullets": ["p1"]}],
        "publications": [{"name": "Paper", "description": "d", "bullets": ["q1"]}]}


def _part_ov(part, **props):
    return _style_overrides({"style": {"parts": {part: props}}}).get("parts", {}).get(part, {})


def test_part_properties_are_validated_and_clamped():
    p = _part_ov("experience.bullets", font="Georgia", size=99, line_height=0.1, color="#AABBCC", bold=True, italic="yes",
                 align="justify", case="upper", letter_spacing=99, margin_top=-5, margin_left=999, padding_bottom=4.5,
                 bullet="–", indent=999, rule_width=2, evil="x")
    assert p == {"font": "Georgia", "size": 36, "line_height": 0.9, "color": "#aabbcc", "bold": True, "align": "justify",
                 "case": "upper", "letter_spacing": 5, "margin_top": 0, "margin_left": 60, "padding_bottom": 4.5,
                 "bullet": "–", "indent": 40}


@pytest.mark.parametrize("bad", [{"color": "red"}, {"color": "#abc"}, {"color": "#12345; x"}, {"align": "start"}, {"case": "title"},
                                 {"bullet": "<b>"}, {"font": "Comic"}, {"bold": 1}, {"margin_top": "4"}])
def test_bad_part_values_are_dropped(bad):
    assert _part_ov("experience.title", **bad) == {}


def test_unknown_parts_and_kind_specific_props_are_dropped():
    assert _style_overrides({"style": {"parts": {"nope.x": {"size": 9}, "experience": {"size": 9}}}}) == {}
    assert _part_ov("experience.title", bullet="•", indent=5, rule_width=2) == {}       # only bullets / headings own these
    assert _part_ov("experience.heading", rule_width=2, rule_color="#112233") == {"rule_width": 2, "rule_color": "#112233"}
    assert _part_ov("experience.heading", rule_width=99) == {"rule_width": 4}


@pytest.mark.parametrize("tpl", PRO)
def test_every_part_selector_matches_something_in_the_template(tpl):
    soup = BeautifulSoup(_render_html(FULL, tpl, "letter", profile_image_path=PNG, profile_image_enabled=True), "html.parser")
    for key, (box, extras) in _STYLE_PARTS.items():
        assert soup.select(box), (tpl, key, box)


def test_part_css_covers_every_property():
    ov = _style_overrides({"style": {"parts": {
        "experience.title": {"font": "Georgia", "size": 12, "color": "#112233", "bold": False, "italic": True, "case": "upper",
                             "letter_spacing": 1, "align": "center", "line_height": 1.5, "margin_top": 3, "padding_left": 4},
        "experience.bullets": {"bullet": "–", "indent": 12},
        "projects.bullets": {"bullet": "none"},
        "experience.heading": {"rule_width": 0}, "skills.heading": {"rule_width": 2, "rule_color": "#445566"},
        "skills.label": {"margin_top": 2}}}})
    css = _style_css(ov)
    assert "'Georgia'" in css and "font-size: 12pt" in css and "color: #112233" in css
    assert "font-weight: 400" in css and "font-style: italic" in css and "text-transform: uppercase" in css
    assert "letter-spacing: 1pt" in css and "text-align: center" in css and "line-height: 1.5" in css
    assert "margin-top: 3pt" in css and "padding-left: 4pt" in css
    assert ".sec-experience ul.bullets li::before { content: '– '" in css and "padding-left: 12pt" in css
    assert ".sec-projects ul.bullets li::before { content: ''" in css
    assert ".sec-experience .section-title { border-bottom: none" in css
    assert "border-bottom: 2pt solid #445566" in css
    assert ".sec-skills .skill-label { margin-top: 2pt; display: inline-block" in css


def test_parts_beat_section_rules_in_the_css():
    css = _style_css(_style_overrides({"style": {"sections": {"experience": {"size": 11}}, "parts": {"experience.meta": {"size": 7}}}}))
    assert css.index("font-size: 10pt") < css.index("font-size: 7pt")       # small lines are the section size - 1, then the part


@pytest.mark.parametrize("tpl", PRO)
def test_select_script_tags_every_part(tpl):
    html = _render_html(FULL, tpl, "letter", select_mode=True)
    assert '"experience.bullets"' in html and '"skills.heading"' in html and "data-part" in html


def _parts_docx(parts, data=FULL):
    st = _docx_style()
    return docx.Document(io.BytesIO(_build_docx({**data, "style": {"parts": parts}}, "letter", st)))


def _para(d, text):
    return next(p for p in d.paragraphs if p.text.startswith(text))


def test_docx_text_properties_per_part():
    d = _parts_docx({"summary.text": {"font": "Georgia", "size": 13, "color": "#112233", "bold": True, "italic": True,
                                       "case": "upper", "letter_spacing": 1.5, "align": "center", "line_height": 1.6},
                     "header.name": {"color": "#aa0000", "size": 30}})
    r = _para(d, "sum").runs[0]
    assert (r.font.name, r.font.size.pt, str(r.font.color.rgb), r.bold, r.italic, r.font.all_caps) == ("Georgia", 13, "112233", True, True, True)
    assert r._r.rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing").get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "30"
    assert _para(d, "sum").alignment == 1 and _para(d, "sum").paragraph_format.line_spacing == 1.6
    name = _para(d, "Dana").runs[0]
    assert str(name.font.color.rgb) == "AA0000" and name.font.size.pt == 30


def test_docx_margin_and_padding_become_spacing_and_indents():
    d = _parts_docx({"experience.description": {"margin_top": 6, "padding_top": 2, "margin_bottom": 3, "margin_left": 10, "padding_left": 5, "margin_right": 4}})
    pf = _para(d, "expdesc").paragraph_format
    assert pf.space_before.pt == 8 and pf.space_after.pt == 3 and pf.left_indent.pt == 15 and pf.right_indent.pt == 4


def test_docx_bullets_char_and_indent():
    d = _parts_docx({"experience.bullets": {"bullet": "–", "indent": 20}, "projects.bullets": {"bullet": "none"}})
    b = _para(d, "– b1")
    assert b.paragraph_format.left_indent.pt == 20
    assert _para(d, "p1").text == "p1"


def test_docx_heading_rule_and_padding():
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    d = _parts_docx({"experience.heading": {"rule_width": 0}, "skills.heading": {"rule_width": 3, "rule_color": "#00ff00", "padding_bottom": 5, "margin_bottom": 7}})
    assert _para(d, "PROFESSIONAL EXPERIENCE")._p.pPr.find(f"{ns}pBdr") is None
    sk = _para(d, "CORE SKILLS")
    bottom = sk._p.pPr.find(f"{ns}pBdr").find(f"{ns}bottom")
    assert (bottom.get(f"{ns}sz"), bottom.get(f"{ns}color"), bottom.get(f"{ns}space")) == ("24", "00ff00", "5")
    assert sk.paragraph_format.space_after.pt == 7


def test_docx_label_and_line_parts_in_skills_and_education():
    d = _parts_docx({"skills.label": {"color": "#0000ff"}, "skills.line": {"size": 12}, "education.school": {"bold": False}, "education.degree": {"italic": True}})
    line = _para(d, "Languages")
    assert str(line.runs[0].font.color.rgb) == "0000FF" and line.runs[1].font.size.pt == 12
    assert _para(d, "TU").runs[0].bold is False and _para(d, "BSc").runs[0].italic is True


def test_docx_ignores_parts_for_unopted_templates():
    off = {k: v for k, v in _docx_style().items() if k != "style_overrides"}
    d = docx.Document(io.BytesIO(_build_docx({**FULL, "style": {"parts": {"summary.text": {"size": 30}}}}, "letter", off)))
    assert _para(d, "sum").runs[0].font.size is None or _para(d, "sum").runs[0].font.size.pt != 30


# Word is strict about the order of pPr children (CT_PPr is an xsd:sequence); python-docx does not check it.
PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
             "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
             "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
             "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u",
             "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath")


def _assert_schema_order(d):
    body = d.element.body
    for kind, order in (("pPr", PPR_ORDER), ("rPr", RPR_ORDER)):
        for el in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + kind):
            tags = [c.tag.split("}")[1] for c in el]
            idx = [order.index(t) for t in tags]
            assert idx == sorted(idx), (kind, tags)


@pytest.mark.parametrize("style", [
    {},
    {"sections": {"experience": {"line_height": 1.4, "font": "Georgia", "size": 11}}},
    {"parts": {"experience.heading": {"margin_top": 6, "margin_left": 3, "align": "center", "line_height": 1.2, "rule_width": 2, "color": "#112233",
                                       "letter_spacing": 1, "case": "upper", "bold": False, "italic": True}}},
    {"parts": {"experience.bullets": {"indent": 12, "margin_left": 4, "align": "right"}, "summary.text": {"letter_spacing": 2, "size": 12, "color": "#445566"},
               "header.contact": {"align": "center", "color": "#101010", "size": 9}, "education.school": {"letter_spacing": 1, "bold": False}}},
])
def test_docx_xml_children_are_in_schema_order(style):
    _assert_schema_order(_docx(style))


# ── section containers and cards: margin and padding ────────────────────────────────────────────────
CARDS = ("experience.entry", "education.entry", "projects.entry", "publications.entry")


def test_section_margin_and_padding_are_validated():
    ov = _style_overrides({"style": {"sections": {"experience": {"size": 10, "margin_top": 99, "padding_left": 4, "margin_right": -2,
                                                                  "margin_bottom": "x", "bogus": 1}}}})
    assert ov["sections"]["experience"] == {"size": 10, "margin_top": 60, "padding_left": 4, "margin_right": 0}


def test_section_box_css_targets_the_container():
    css = _style_css(_style_overrides({"style": {"sections": {"experience": {"margin_top": 5, "padding_left": 3}, "header": {"padding_bottom": 2}}}}))
    assert ".sec-experience { margin-top: 5pt; padding-left: 3pt; }" in css
    assert ".header { padding-bottom: 2pt; }" in css


@pytest.mark.parametrize("card", CARDS)
def test_cards_are_parts_with_the_full_property_set(card):
    box = _STYLE_PARTS[card][0]
    css = _style_css(_style_overrides({"style": {"parts": {card: {"margin_bottom": 4, "padding_top": 2, "color": "#112233"}}}}))
    assert f"{box} {{ margin-bottom: 4pt; padding-top: 2pt; }}" in css and f"{box} {{ color: #112233; }}" in css


TWO_JOBS = {**FULL, "experience": [dict(FULL["experience"][0], title="PM1", bullets=["a1"]), dict(FULL["experience"][0], title="PM2", bullets=["a2"])]}


def _style_docx(style, data=FULL):
    return docx.Document(io.BytesIO(_build_docx({**data, "style": style}, "letter", _docx_style())))


def test_docx_section_box_spaces_and_indents_the_whole_section():
    d = _style_docx({"sections": {"experience": {"margin_top": 10, "padding_top": 2, "margin_bottom": 7, "margin_left": 8, "padding_left": 1}}})
    heading, title, bullet = _para(d, "PROFESSIONAL EXPERIENCE"), _para(d, "PM - Acme"), _para(d, "• b1")
    assert heading.paragraph_format.space_before.pt == pytest.approx(4.2 + 12)
    assert bullet.paragraph_format.space_after.pt == pytest.approx(0.45 + 7)
    assert title.paragraph_format.left_indent.pt == pytest.approx(9) and bullet.paragraph_format.left_indent.pt == pytest.approx(9.7 + 9)
    assert _para(d, "Languages").paragraph_format.left_indent is None              # other sections untouched


def test_docx_header_box_wraps_the_name_block():
    d = _style_docx({"sections": {"header": {"margin_left": 6}}})
    assert _para(d, "Dana").paragraph_format.left_indent.pt == pytest.approx(6)
    assert _para(d, "PM").paragraph_format.left_indent.pt == pytest.approx(6)
    assert _para(d, "PROFESSIONAL SUMMARY").paragraph_format.left_indent is None


def test_docx_cards_space_each_entry_on_its_own():
    d = _parts_docx({"experience.entry": {"margin_top": 6, "margin_bottom": 5, "margin_left": 4, "color": "#112233"}}, data=TWO_JOBS)
    t1, b1, t2, b2 = _para(d, "PM1"), _para(d, "• a1"), _para(d, "PM2"), _para(d, "• a2")
    assert t1.paragraph_format.space_before.pt == pytest.approx(3 + 6) and t2.paragraph_format.space_before.pt == pytest.approx(3 + 6)
    assert b1.paragraph_format.space_after.pt == pytest.approx(0.45 + 5) and b2.paragraph_format.space_after.pt == pytest.approx(0.45 + 5)
    assert t1.paragraph_format.left_indent.pt == pytest.approx(4) and b2.paragraph_format.left_indent.pt == pytest.approx(9.7 + 4)
    assert str(t1.runs[0].font.color.rgb) == "112233"                                 # a card's text properties reach its runs
    assert _para(d, "Languages").paragraph_format.left_indent is None


def test_docx_education_card_indents_school_and_degree():
    d = _parts_docx({"education.entry": {"margin_left": 4}})
    assert _para(d, "TU").paragraph_format.left_indent.pt == pytest.approx(4) and _para(d, "BSc").paragraph_format.left_indent.pt == pytest.approx(4)


@pytest.mark.parametrize("style", [
    {"sections": {"experience": {"margin_top": 4, "padding_bottom": 3, "margin_left": 2, "line_height": 1.3}, "header": {"margin_left": 5, "margin_top": 2}}},
    {"parts": {c: {"margin_top": 3, "margin_left": 2, "padding_bottom": 1, "align": "center"} for c in CARDS}},
])
def test_docx_section_and_card_xml_stay_in_schema_order(style):
    _assert_schema_order(docx.Document(io.BytesIO(_build_docx({**TWO_JOBS, "style": style}, "letter", _docx_style()))))


# ── Manrope: a font that ships with the templates (fonts/ folder), embedded only when chosen ─────────
from backend.api.routes_resumes import _STYLE_FONTS

FONT_FILES = sorted((ROOT / "professional-top-left-image" / "fonts").glob("Manrope-*.ttf"))


def test_manrope_is_a_choosable_font():
    assert _STYLE_FONTS["Manrope"].startswith("'Manrope'")
    ov = _style_overrides({"style": {"sections": {"summary": {"font": "Manrope"}}, "parts": {"summary.text": {"font": "Manrope"}}}})
    assert ov["sections"]["summary"] == {"font": "Manrope"} and ov["parts"]["summary.text"] == {"font": "Manrope"}


def test_font_face_rules_cover_only_the_fonts_in_use_and_known_weights():
    fonts = {"Manrope-Bold.ttf": "data:font/truetype;base64,AAA", "Manrope-Light.ttf": "data:font/truetype;base64,BBB",
             "Manrope-Weird.ttf": "data:x", "Other-Regular.ttf": "data:y"}
    css = _style_css(_style_overrides({"style": {"parts": {"summary.text": {"font": "Manrope"}}}}), fonts=fonts)
    assert css.count("@font-face") == 2
    assert "font-weight: 700" in css and "font-weight: 300" in css and "AAA" in css and "BBB" in css
    assert "Weird" not in css and "data:x" not in css and "data:y" not in css
    assert "@font-face" not in _style_css(_style_overrides({"style": {"parts": {"summary.text": {"font": "Georgia"}}}}), fonts=fonts)
    assert "@font-face" not in _style_css(_style_overrides({"style": {"parts": {"summary.text": {"size": 9}}}}), fonts=fonts)


@pytest.mark.skipif(not FONT_FILES, reason="the Manrope .ttf files are not checked out (they are git-ignored)")
@pytest.mark.parametrize("tpl", PRO)
def test_manrope_is_embedded_in_the_pdf_html_only_when_used(tpl):
    assert "@font-face" not in _html(tpl)
    used = _html(tpl, {"sections": {"experience": {"font": "Manrope"}}})
    assert used.count("@font-face") == len(FONT_FILES) == 7 and "data:font/truetype;base64," in used
    assert all(f"font-weight: {w}" in used for w in (200, 300, 400, 500, 600, 700, 800))
    assert "font-family: 'Manrope','Arial',sans-serif" in used
    assert "@font-face" not in _html(tpl, {"sections": {"experience": {"font": "Georgia"}}})


def test_docx_names_manrope_in_the_runs():
    d = _parts_docx({"summary.text": {"font": "Manrope"}})
    assert _para(d, "sum").runs[0].font.name == "Manrope"
