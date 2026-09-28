"""`json_data["section_titles"]` overrides the section headings in every template and in Word exports;
blank or missing values keep the template's own heading."""
import io
import json
from pathlib import Path

import docx
import pytest

from backend.api.routes_persona import _resume_content_from
from backend.api.routes_resumes import _build_docx, _render_html

ROOT = Path(__file__).resolve().parents[1] / "resume_templates"
TEMPLATES = sorted(p.name for p in ROOT.iterdir() if p.is_dir())
DATA = {"header": {"name": "Dana", "contact_items": []}, "summary": "s",
        "experience": [{"title": "PM", "company": "Acme", "bullets": ["b"]}],
        "skills": {"Languages": "Go"}, "education": [{"school": "TU", "degree": "BSc"}],
        "projects": [{"name": "Side"}], "publications": [{"name": "Paper"}]}
KEYS = ("summary", "skills", "experience", "education", "projects", "publications")
CUSTOM = {k: f"Custom {k} <&>" for k in KEYS}


def _titles(html):
    return [t.strip() for t in html.split('<div class="section-title">')[1:]]


@pytest.mark.parametrize("template", TEMPLATES)
def test_custom_titles_replace_headings_and_are_escaped(template):
    html = _render_html({**DATA, "section_titles": CUSTOM}, template, "a4")
    for k in KEYS:
        assert f"Custom {k} &lt;&amp;&gt;</div>" in html
    assert "Custom summary <&>" not in html
    assert len(_titles(html)) == 6


@pytest.mark.parametrize("template", TEMPLATES)
@pytest.mark.parametrize("titles", [None, {}, {k: "  " for k in KEYS}, {"experience": 5}, "junk"])
def test_blank_titles_keep_template_default(template, titles):
    base = _titles(_render_html(DATA, template, "a4"))
    assert _titles(_render_html({**DATA, "section_titles": titles}, template, "a4")) == base
    assert 5 <= len(base) <= 6      # only some templates print a Summary heading by default


def test_docx_custom_titles_styled_and_plain():
    style = json.loads((ROOT / "professional-top-left-image" / "meta.json").read_text())["docx"]
    for st in (style, None):
        d = docx.Document(io.BytesIO(_build_docx({**DATA, "section_titles": CUSTOM}, "a4", st)))
        text = [p.text for p in d.paragraphs]
        assert all(f"Custom {k} <&>" in text for k in KEYS)


def test_docx_blank_titles_keep_defaults():
    style = json.loads((ROOT / "professional-top-left-image" / "meta.json").read_text())["docx"]
    d = docx.Document(io.BytesIO(_build_docx({**DATA, "section_titles": {"summary": ""}}, "a4", style)))
    assert "PROFESSIONAL SUMMARY" in [p.text for p in d.paragraphs]
    plain = docx.Document(io.BytesIO(_build_docx({**DATA, "section_titles": {"summary": ""}}, "a4")))
    assert "Summary" not in [p.text for p in plain.paragraphs]


def test_persona_import_keeps_section_titles():
    assert _resume_content_from({"section_titles": {"skills": "Tech"}})["section_titles"] == {"skills": "Tech"}
