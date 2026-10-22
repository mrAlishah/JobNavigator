"""Regenerate the fork's golden files from the current code.

    python -m backend.tests.fixtures.generate_fork_golden

`fork_golden_html.json`   sha256 and length of every template's rendered HTML (base64 data replaced by a
                          length + hash marker), plain and fully styled
`fork_golden_docx.json`   the structure of the Word export per template (text, run and paragraph properties,
                          links, page setup, header/footer, image sizes, canonical-XML hashes), not zip bytes
`fork_golden_routes.json` every route's methods, path, name and dependencies, plus a hash of the OpenAPI schema

`test_fork_golden.py` compares the code against these files. Run this only after an intended change (an
upstream merge, a refactor), then review the diff of the three files before committing them.
"""
import hashlib
import io
import json
import pathlib
import re

import docx
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from backend.api import routes_resumes as rr

HERE = pathlib.Path(__file__).parent
FILES = {name: HERE / f"fork_golden_{name}.json" for name in ("html", "docx", "routes")}

BASE = {
    "header": {"name": "Dana Okonkwo", "title": "Senior Product Manager",
               "contact_items": [{"text": "dana@example.com", "url": "mailto:dana@example.com"},
                                 {"text": "+1 555 0100", "url": "tel:+15550100"},
                                 {"text": "Berlin", "url": ""}]},
    "summary": "Product leader with **ten** years in platforms.",
    "experience": [
        {"title": "PM", "company": "Acme", "location": "Berlin", "date": "2020 - 2024",
         "description": "Ran the platform.", "bullets": ["Shipped **Go** services", "Cut cost 20%"]},
        {"title": "APM", "company": "Globex", "location": "Remote", "date": "2017 - 2020",
         "description": "", "bullets": ["Owned onboarding"]},
    ],
    "skills": {"Languages": "Go, Python", "Tools": "Jira, SQL"},
    "education": [{"school": "TU Berlin", "location": "Berlin", "degree": "BSc CS"}],
    "projects": [{"name": "Side", "url": "example.com/side", "description": "A thing", "bullets": ["Did it"]}],
    "publications": [{"name": "A paper", "description": "Journal", "bullets": ["Cited twice"]}],
}
FULL = {
    **BASE,
    "section_titles": {"summary": "Profile", "experience": "Career", "skills": "Stack"},
    "style": {
        "page": {"top": 0.6, "right": 0.7, "bottom": 0.6, "left": 0.7},
        "image": {"width": 120, "height": 90},
        "sections": {"summary": {"font": "Georgia", "size": 11, "margin_top": 4, "padding_left": 3},
                     "header": {"line_height": 1.4}},
        "parts": {
            "header.name": {"font": "Manrope", "size": 22, "color": "#112233", "bold": True},
            "experience.title": {"bold": True, "case": "upper", "letter_spacing": 0.5},
            "experience.bullets": {"bullet": "–", "indent": 14},
            "experience.entry": {"margin_bottom": 6, "padding_top": 2},
            "summary.heading": {"rule_color": "#ff0000", "rule_width": 2, "align": "center"},
            "skills.label": {"italic": True},
        },
    },
}
FOOTER = "Dana Okonkwo - Senior Product Manager - Acme"
PNG_URI = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def _templates():
    return sorted(rr._discover_templates(), key=lambda t: t["id"])


def _norm_html(html: str) -> str:
    return re.sub(r"data:[^\"')\s>]+", lambda m: f"data:<{len(m.group())}:{hashlib.sha256(m.group().encode()).hexdigest()[:12]}>", html)


def html_snapshot() -> dict:
    out = {}
    for t in _templates():
        tid, opts = t["id"], bool((t.get("docx") or {}).get("style_overrides"))
        cases = [("plain|letter", BASE, "letter", 0, None, True, False),
                 ("full|a4", FULL, "a4", 0.5, PNG_URI, True, False),
                 ("no-image|letter", FULL, "letter", 0, PNG_URI, False, False)]
        if opts:
            cases.append(("select|letter", FULL, "letter", 0.5, PNG_URI, True, True))
        for label, data, fmt, reserve, image, enabled, select in cases:
            html = _norm_html(rr._render_html(data, tid, fmt, reserve, image, enabled, select))
            out[f"{tid}|{label}"] = {"sha256": hashlib.sha256(html.encode()).hexdigest(), "len": len(html)}
    return out


def _canon(el) -> str:
    return re.sub(r' xmlns(:\w+)?="[^"]*"', "", etree.tostring(el, method="c14n").decode())


def _png(w, h) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (180, 40, 40)).save(buf, "PNG")
    return buf.getvalue()


def docx_structure(blob: bytes) -> dict:
    d = docx.Document(io.BytesIO(blob))
    paras = []
    for p in d.element.body.iter(qn("w:p")):
        ppr = p.find(qn("w:pPr"))
        paras.append({
            "text": "".join(t.text or "" for t in p.iter(qn("w:t"))),
            "pPr": _canon(ppr) if ppr is not None else "",
            "runs": [{"text": "".join(t.text or "" for t in r.iter(qn("w:t"))),
                      "rPr": _canon(r.find(qn("w:rPr"))) if r.find(qn("w:rPr")) is not None else ""}
                     for r in p.iter(qn("w:r"))],
        })
    links, images = [], []
    for rel in d.part.rels.values():
        if rel.reltype.endswith("/hyperlink"):
            links.append(rel.target_ref)
        elif rel.reltype.endswith("/image"):
            with Image.open(io.BytesIO(rel.target_part.blob)) as im:
                images.append(list(im.size))
    sec = d.sections[0]
    return {
        "paragraphs": paras,
        "tables": len(d.tables),
        "hyperlinks": sorted(links),
        "image_pixels": sorted(images),
        "inline_shapes": [[s.width, s.height] for s in d.inline_shapes],
        "page": [sec.page_width, sec.page_height, sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin],
        "header": [p.text for p in sec.header.paragraphs],
        "footer": [p.text for p in sec.footer.paragraphs],
        "document_sha256": hashlib.sha256(_canon(d.element).encode()).hexdigest(),
        "styles_sha256": hashlib.sha256(_canon(d.styles.element).encode()).hexdigest(),
    }


def docx_snapshot() -> dict:
    out = {}
    for t in _templates():
        style = t.get("docx")
        for label, data, footer, image in (("plain", BASE, "", None),
                                           ("full", FULL, FOOTER, _png(60, 40) if style and style.get("image") else None)):
            blob = rr._build_docx(rr._xml_safe(data), "letter", style, image, rr._xml_safe(footer))
            out[f"{t['id']}|{label}"] = docx_structure(blob)
    return out


def routes_snapshot() -> dict:
    from backend.main import app

    def deps(dependant):
        names = []
        for sub in dependant.dependencies:
            names.append(getattr(sub.call, "__name__", repr(sub.call)))
            names += deps(sub)
        return names

    routes = sorted(
        ({"methods": sorted(r.methods or []), "path": r.path, "name": r.name, "deps": sorted(set(deps(r.dependant)))}
         for r in app.routes if hasattr(r, "dependant")),
        key=lambda r: (r["path"], r["methods"]))
    schema = json.dumps(app.openapi(), sort_keys=True)
    return {"routes": routes, "openapi_sha256": hashlib.sha256(schema.encode()).hexdigest()}


SNAPSHOTS = {"html": html_snapshot, "docx": docx_snapshot, "routes": routes_snapshot}


def main() -> None:
    for name, path in FILES.items():
        path.write_text(json.dumps(SNAPSHOTS[name](), indent=1, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
        print("wrote %s" % path)


if __name__ == "__main__":
    main()
